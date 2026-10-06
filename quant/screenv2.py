"""
AUTO SCREENER UNION V2 - versi final.
Rule: A + T + V (union)
Exit: TP di high aktual kalau high >= +3%, SL di -2%, else PROFIT/CLOSE.
Filter bulan: berdasarkan tanggal EKSEKUSI (T+1).
"""

import os
import subprocess
import sys
from datetime import datetime
import pandas as pd
import numpy as np

from .data_loader import load_panel_from_gabungan
from .data_cleaner import clean_panel
from .feature_engineering import build_features
from .event_detector import detect_events


COST_RT = 0.008
TP_THRESHOLD = 0.03
SL_PCT = -0.02


def get_bulan_berjalan():
    today = pd.Timestamp(datetime.now().date())
    start = today.replace(day=1)
    if today.month == 12:
        end = today.replace(year=today.year + 1, month=1, day=1) - pd.Timedelta(days=1)
    else:
        end = today.replace(month=today.month + 1, day=1) - pd.Timedelta(days=1)
    return start, end, today.strftime("%B %Y"), today.strftime("%Y-%m")


def jalankan_update():
    print(f"\n{'='*100}")
    print("STEP 1: Update data dari IDX")
    print(f"{'='*100}\n")

    today = datetime.now()
    if today.weekday() >= 5:
        print("[SKIP] Hari ini weekend, tidak ada data baru.")
        return True

    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    update_script = os.path.join(root, "update.py")
    gabung_script = os.path.join(root, "gabung.py")

    if not os.path.exists(update_script):
        print(f"[!] File {update_script} tidak ditemukan.")
        return False

    print("Menjalankan update.py...")
    try:
        result = subprocess.run([sys.executable, update_script],
                                 cwd=root, timeout=1800)
        if result.returncode != 0:
            print(f"[!] Error update (code {result.returncode})")
            return False
    except subprocess.TimeoutExpired:
        print("[!] Update timeout. Lanjut.")
        return False
    except Exception as e:
        print(f"[!] Gagal update: {e}")
        return False

    if os.path.exists(gabung_script):
        print("\nMenjalankan gabung.py...")
        try:
            subprocess.run([sys.executable, gabung_script], cwd=root, timeout=300)
        except Exception:
            pass

    print("\n[OK] Update selesai.")
    return True


def rule_a(d):
    return ((d["ret_1d"] > 0.15) & (d["ret_5d"] > 0.15) &
            (d["px_vs_sma20"] > 0.20))


def rule_t(d):
    return ((d["ret_1d"] > 0.15) & (d["px_vs_sma20"] > 0.20) &
            (d["rk_turnover"] > 0.90))


def rule_v(d):
    return ((d["ret_5d"] > 0.15) & (d["rk_turnover"] > 0.90) &
            (d["vol_regime_high"] == 1))


def rule_union(d):
    return rule_a(d) | rule_t(d) | rule_v(d)


def next_trading_day(ts):
    nxt = pd.Timestamp(ts) + pd.Timedelta(days=1)
    while nxt.weekday() >= 5:
        nxt += pd.Timedelta(days=1)
    return nxt


def add_features(df):
    df = df.copy()
    df = df.sort_values(["kode", "date"]).reset_index(drop=True)

    if "listed_shares" in df.columns:
        df["turnover"] = df["volume"] / df["listed_shares"].replace(0, np.nan)
        df["rk_turnover"] = df.groupby("date")["turnover"].rank(pct=True)
    else:
        df["rk_turnover"] = 0

    df["vol_regime_high"] = (df["atr_pct_rank_60"] > 0.8).astype(int)

    for h in [1, 2, 3]:
        df[f"open_T{h}"] = df.groupby("kode")["open"].shift(-h)
        df[f"high_T{h}"] = df.groupby("kode")["high"].shift(-h)
        df[f"low_T{h}"] = df.groupby("kode")["low"].shift(-h)
        df[f"close_T{h}"] = df.groupby("kode")["close"].shift(-h)

    return df


def simulate(row):
    """
    Exit logic:
    1. High >= entry x 1.03 : TP, exit di HIGH aktual
    2. Low <= entry x 0.98  : SL, exit di -2%
    3. High > entry         : PROFIT, exit di HIGH
    4. Sisanya              : CLOSE
    """
    close_t = row["close"]
    entry = row["open_T1"]

    if pd.isna(entry) or entry <= 0:
        return None, "NO_DATA"
    if entry < close_t:
        return None, "SKIP_GAP_DOWN"

    high = row["high_T1"]
    low = row["low_T1"]
    close_t1 = row["close_T1"]

    if pd.isna(high) or high <= 0:
        return None, "NO_DATA"

    # TP
    if high >= entry * (1 + TP_THRESHOLD):
        return high / entry - 1 - COST_RT, "TP"

    # SL
    if not pd.isna(low) and low <= entry * (1 + SL_PCT):
        return SL_PCT - COST_RT, "SL"

    # PROFIT
    if high > entry:
        return high / entry - 1 - COST_RT, "PROFIT"

    # CLOSE
    if pd.isna(close_t1) or close_t1 <= 0:
        return None, "NO_DATA"
    return close_t1 / entry - 1 - COST_RT, "CLOSE"


def excel_style(ws, title_text, subtitle_text, header_row=4):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=ws.max_column)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title_text
    title_cell.font = Font(bold=True, size=16, color="1F4E78")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=ws.max_column)
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = subtitle_text
    sub_cell.font = Font(italic=True, size=10, color="666666")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 15

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78",
                               fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 30

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate


def autofit_columns(ws, skip_rows=3):
    from openpyxl.utils import get_column_letter
    for col_idx, col in enumerate(ws.iter_cols(min_row=skip_rows + 1,
                                                 max_row=ws.max_row), 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 22)


def simpan_excel(kandidat, signals_bulan, pred_date, path_dir,
                 bulan_label, bulan_file):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule
    except ImportError:
        print("  [!] openpyxl tidak terinstall.")
        return

    base_name = f"Rekap_Union_{bulan_file}"
    filepath = os.path.join(path_dir, f"{base_name}.xlsx")

    if os.path.exists(filepath):
        try:
            with open(filepath, 'a'):
                pass
        except PermissionError:
            ts = datetime.now().strftime("%H%M%S")
            filepath = os.path.join(path_dir, f"{base_name}_{ts}.xlsx")
            print(f"  [!] File lama sedang dibuka, disimpan sebagai: {filepath}")

    # SHEET 1: KANDIDAT
    if not kandidat.empty:
        df_kand = kandidat[["kode", "close", "value", "ret_1d", "ret_5d",
                             "px_vs_sma20", "rk_turnover", "vol_regime_high",
                             "match", "tanggal_eksekusi"]].copy()
        df_kand["ret_1d"] = (df_kand["ret_1d"] * 100).round(2)
        df_kand["ret_5d"] = (df_kand["ret_5d"] * 100).round(2)
        df_kand["px_vs_sma20"] = (df_kand["px_vs_sma20"] * 100).round(2)
        df_kand["rk_turnover"] = df_kand["rk_turnover"].round(3)
        df_kand["value"] = (df_kand["value"] / 1e9).round(2)
        df_kand = df_kand.rename(columns={
            "kode": "Kode", "close": "Close T", "value": "Value (M)",
            "ret_1d": "Ret1d%", "ret_5d": "Ret5d%",
            "px_vs_sma20": "Px>sma20%", "rk_turnover": "Turnover",
            "vol_regime_high": "VolHigh", "match": "Match",
            "tanggal_eksekusi": "Eksekusi",
        })
        cols = ["Eksekusi", "Kode", "Close T", "Value (M)", "Ret1d%",
                "Ret5d%", "Px>sma20%", "Turnover", "VolHigh", "Match"]
        df_kand = df_kand[cols]
        df_kand.insert(0, "No", range(1, len(df_kand) + 1))
    else:
        df_kand = pd.DataFrame({"Info": ["Tidak ada kandidat"]})

    # SHEET 2: REVIEW HARIAN
    if not signals_bulan.empty:
        review_rows = []
        for _, row in signals_bulan.iterrows():
            tgl_sinyal = pd.Timestamp(row["date"])
            tgl_eksekusi = next_trading_day(tgl_sinyal)
            ret, status = simulate(row)

            match = ""
            if row.get("sig_a", False):
                match += "A"
            if row.get("sig_t", False):
                match += "T"
            if row.get("sig_v", False):
                match += "V"

            entry = row["open_T1"]
            high = row["high_T1"]

            review_rows.append({
                "Tanggal": tgl_eksekusi.date(),
                "Hari": ['Senin', 'Selasa', 'Rabu', 'Kamis', 'Jumat',
                         'Sabtu', 'Minggu'][tgl_eksekusi.weekday()],
                "Sinyal Dari": tgl_sinyal.date(),
                "Kode": row["kode"],
                "Close T": round(row["close"], 0),
                "Entry": round(entry, 0) if not pd.isna(entry) else "",
                "High T+1": round(high, 0) if not pd.isna(high) else "",
                "Match": match,
                "Status": status,
                "P/L %": round(ret * 100, 2) if ret is not None else "",
            })
        df_review = pd.DataFrame(review_rows)
        df_review = df_review.sort_values(["Tanggal", "Kode"]).reset_index(drop=True)
    else:
        df_review = pd.DataFrame({
            "Info": [f"Belum ada eksekusi di {bulan_label}"]
        })

    # SHEET 3: REVIEW MINGGUAN
    minggu_rows = []
    if not signals_bulan.empty:
        sr = signals_bulan.copy()
        sr["tanggal_sinyal"] = pd.to_datetime(sr["date"])
        sr["tanggal_eksekusi"] = sr["tanggal_sinyal"].apply(next_trading_day)

        if not sr.empty:
            sr["dom"] = sr["tanggal_eksekusi"].dt.day
            sr["minggu_ke"] = ((sr["dom"] - 1) // 7) + 1

            for wk in sorted(sr["minggu_ke"].unique()):
                wk_data = sr[sr["minggu_ke"] == wk]
                n_sig = len(wk_data)
                n_exec = 0
                n_win = 0
                n_tp = 0
                n_sl = 0
                n_profit = 0
                n_close = 0
                n_skip = 0
                n_pending = 0
                pl_total = 0

                for _, row in wk_data.iterrows():
                    ret, status = simulate(row)
                    if status in ("TP", "PROFIT", "SL", "CLOSE"):
                        n_exec += 1
                        pl_total += ret
                        if ret > 0:
                            n_win += 1
                        if status == "TP":
                            n_tp += 1
                        elif status == "PROFIT":
                            n_profit += 1
                        elif status == "SL":
                            n_sl += 1
                        else:
                            n_close += 1
                    elif status == "SKIP_GAP_DOWN":
                        n_skip += 1
                    else:
                        n_pending += 1

                tgl_min = wk_data["tanggal_eksekusi"].min().date()
                tgl_max = wk_data["tanggal_eksekusi"].max().date()

                minggu_rows.append({
                    "Minggu Ke": f"Minggu {wk}",
                    "Periode": f"{tgl_min} s/d {tgl_max}",
                    "Sinyal": n_sig,
                    "Eksekusi": n_exec,
                    "Profit": n_win,
                    "TP": n_tp,
                    "SL": n_sl,
                    "Skip": n_skip,
                    "Pending": n_pending,
                    "Win Rate %": round(n_win / n_exec * 100, 1) if n_exec > 0 else 0,
                    "Total P/L %": round(pl_total * 100, 2),
                    "Avg P/L %": round(pl_total / n_exec * 100, 2) if n_exec > 0 else 0,
                })

    if minggu_rows:
        df_minggu = pd.DataFrame(minggu_rows)
    else:
        df_minggu = pd.DataFrame({
            "Info": [f"Belum ada eksekusi di {bulan_label}"]
        })

    # SHEET 4: RINGKASAN
    if not signals_bulan.empty:
        total_pl = 0
        n_exec = 0
        n_win = 0
        n_tp = 0
        n_sl = 0
        n_profit = 0
        n_close = 0
        n_skip = 0
        n_pending = 0

        for _, row in signals_bulan.iterrows():
            ret, status = simulate(row)
            if status in ("TP", "PROFIT", "SL", "CLOSE"):
                n_exec += 1
                total_pl += ret
                if ret > 0:
                    n_win += 1
                if status == "TP":
                    n_tp += 1
                elif status == "PROFIT":
                    n_profit += 1
                elif status == "SL":
                    n_sl += 1
                else:
                    n_close += 1
            elif status == "SKIP_GAP_DOWN":
                n_skip += 1
            else:
                n_pending += 1

        win_rate = (n_win / n_exec * 100) if n_exec > 0 else 0
        avg_pl = (total_pl / n_exec * 100) if n_exec > 0 else 0

        n_a = int(signals_bulan["sig_a"].sum())
        n_t = int(signals_bulan["sig_t"].sum())
        n_v = int(signals_bulan["sig_v"].sum())

        ringkasan_rows = [
            ["PERIODE EKSEKUSI", f"Bulan {bulan_label}"],
            ["Exit Rule", f"TP +{TP_THRESHOLD*100:.0f}% | SL {SL_PCT*100:.0f}%"],
            ["", ""],
            ["RINGKASAN SINYAL", ""],
            ["Total Sinyal", len(signals_bulan)],
            ["  Rule A (momentum)", n_a],
            ["  Rule T (turnover)", n_t],
            ["  Rule V (volatility)", n_v],
            ["", ""],
            ["EKSEKUSI", ""],
            ["Tereksekusi", n_exec],
            ["  TP (high >= +3%)", n_tp],
            ["  SL (kena -2%)", n_sl],
            ["  Skip (gap down)", n_skip],
            ["  Pending (belum T+1)", n_pending],
            ["", ""],
            ["PERFORMA", ""],
            ["Win Rate", f"{win_rate:.1f}%"],
            ["Total P/L", f"{total_pl*100:+.2f}%"],
            ["Avg P/L per Trade", f"{avg_pl:+.2f}%"],
            ["", ""],
            ["VERDICT", ""],
            ["Status", "[OK] LAYAK" if win_rate >= 65 and avg_pl > 2
             else ("[!] MARGINAL" if win_rate >= 50 else "[X] EVALUASI")],
        ]
    else:
        ringkasan_rows = [
            ["PERIODE EKSEKUSI", f"Bulan {bulan_label}"],
            ["Info", "Belum ada sinyal"],
        ]

    df_ringkasan = pd.DataFrame(ringkasan_rows, columns=["Metrik", "Nilai"])

    # TULIS EXCEL
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df_ringkasan.to_excel(writer, sheet_name="Ringkasan",
                               index=False, startrow=3)
        df_kand.to_excel(writer, sheet_name="Kandidat Besok",
                          index=False, startrow=3)
        df_minggu.to_excel(writer, sheet_name="Review Mingguan",
                            index=False, startrow=3)
        df_review.to_excel(writer, sheet_name="Review Harian",
                            index=False, startrow=3)

    # FORMATTING
    wb = load_workbook(filepath)

    GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    GREEN_FONT = Font(color="006100", bold=True)
    RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    RED_FONT = Font(color="9C0006", bold=True)
    SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    SECTION_FONT = Font(bold=True, size=11, color="1F4E78")

    ws = wb["Ringkasan"]
    excel_style(ws, f"RINGKASAN UNION - {bulan_label.upper()}",
                f"Dibuat: {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
                f"Tanggal = hari eksekusi (T+1)")
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=2):
        cell = row[0]
        if cell.value and isinstance(cell.value, str):
            if cell.value.isupper() and cell.value.strip():
                cell.fill = SECTION_FILL
                cell.font = SECTION_FONT
                row[1].fill = SECTION_FILL
    autofit_columns(ws)

    ws = wb["Kandidat Besok"]
    excel_style(ws, f"KANDIDAT EKSEKUSI {pred_date.strftime('%A, %d %B %Y').upper()}",
                f"Sinyal dari close hari sebelumnya | Entry open {pred_date}")
    autofit_columns(ws)
    if not kandidat.empty:
        for row in ws.iter_rows(min_row=5, max_row=ws.max_row,
                                 min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")

    ws = wb["Review Mingguan"]
    excel_style(ws, f"REVIEW MINGGUAN - {bulan_label.upper()}",
                f"Berdasarkan tanggal EKSEKUSI")
    autofit_columns(ws)
    if minggu_rows:
        for i, cell in enumerate(ws[4], 1):
            if cell.value in ("Total P/L %", "Avg P/L %"):
                col_letter = get_column_letter(i)
                rng = f"{col_letter}5:{col_letter}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="greaterThan", formula=["0"],
                               fill=GREEN_FILL, font=GREEN_FONT))
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="lessThan", formula=["0"],
                               fill=RED_FILL, font=RED_FONT))

    ws = wb["Review Harian"]
    excel_style(ws, f"REVIEW HARIAN - {bulan_label.upper()}",
                f"Tanggal = hari EKSEKUSI (T+1) | 'Sinyal Dari' = hari sinyal muncul")
    autofit_columns(ws)
    if not df_review.empty and "P/L %" in df_review.columns:
        for i, cell in enumerate(ws[4], 1):
            if cell.value == "P/L %":
                col_letter = get_column_letter(i)
                rng = f"{col_letter}5:{col_letter}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="greaterThan", formula=["0"],
                               fill=GREEN_FILL, font=GREEN_FONT))
                ws.conditional_formatting.add(rng,
                    CellIsRule(operator="lessThan", formula=["0"],
                               fill=RED_FILL, font=RED_FONT))

    wb.active = wb.sheetnames.index("Kandidat Besok")
    wb.save(filepath)
    print(f"  -> Excel: {filepath}")


def main():
    print("=" * 100)
    print(f"AUTO SCREENER UNION V2 - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"TP +{TP_THRESHOLD*100:.0f}% | SL {SL_PCT*100:.0f}%")
    print("=" * 100)

    bulan_start, bulan_end, bulan_label, bulan_file = get_bulan_berjalan()
    print(f"\nBulan berjalan : {bulan_label}")
    print(f"Periode        : {bulan_start.date()} s/d {bulan_end.date()}")

    jalankan_update()

    print(f"\n{'='*100}")
    print("STEP 2: Proses screening")
    print(f"{'='*100}\n")

    print("Memuat data...")
    try:
        panel = load_panel_from_gabungan()
    except FileNotFoundError:
        print("[!] File data tidak ditemukan.")
        return

    panel = clean_panel(panel)
    panel = build_features(panel)
    panel = detect_events(panel)
    panel = add_features(panel)
    panel = panel[(panel["open"] > 0) & (panel["close"] > 0)].copy()

    last_date = pd.Timestamp(panel["date"].max())
    today = pd.Timestamp(datetime.now().date())
    pred_date = next_trading_day(last_date)

    print(f"\nData terakhir    : {last_date.date()} "
          f"({['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu'][last_date.weekday()]})")
    print(f"Tanggal eksekusi : {pred_date.date()} "
          f"({['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu'][pred_date.weekday()]})")

    # Kandidat
    today_data = panel[panel["date"] == last_date].copy()
    m_a = rule_a(today_data).fillna(False).values
    m_t = rule_t(today_data).fillna(False).values
    m_v = rule_v(today_data).fillna(False).values
    m_union = m_a | m_t | m_v

    kandidat = today_data[m_union].copy()

    print(f"\n{'#'*100}")
    print(f"#  KANDIDAT EKSEKUSI {pred_date.date()}")
    print(f"#  (Sinyal dari close {last_date.date()})")
    print(f"{'#'*100}")

    if not kandidat.empty:
        kandidat["match"] = ""
        kandidat.loc[m_a[m_union], "match"] += "A"
        kandidat.loc[m_t[m_union], "match"] += "T"
        kandidat.loc[m_v[m_union], "match"] += "V"
        kandidat["tanggal_eksekusi"] = pred_date.date()
        kandidat = kandidat.sort_values("value", ascending=False)

        print(f"\n  {len(kandidat)} KANDIDAT:\n")
        print(f"  {'No':>3} {'Kode':<7} {'Close T':>9} {'Ret1d%':>8} "
              f"{'Ret5d%':>8} {'Px>sma20':>9} {'Turnover':>9} "
              f"{'Match':>6} {'Value(M)':>10}")
        print("  " + "-" * 95)

        for i, (_, r) in enumerate(kandidat.iterrows(), 1):
            turnover = f"{r['rk_turnover']:.2f}" if not pd.isna(r['rk_turnover']) else "n/a"
            print(f"  {i:>3} {r['kode']:<7} {r['close']:>9.0f} "
                  f"{r['ret_1d']*100:>+7.1f}% {r['ret_5d']*100:>+7.1f}% "
                  f"{r['px_vs_sma20']*100:>+8.1f}% {turnover:>9} "
                  f"{r['match']:>6} {r['value']/1e9:>9.2f}")
    else:
        print(f"\n  Tidak ada kandidat untuk {pred_date.date()}.")

    # Review bulan berjalan
    print(f"\n{'#'*100}")
    print(f"#  REVIEW BULAN {bulan_label.upper()}")
    print(f"#  (Filter: tanggal eksekusi di bulan ini)")
    print(f"{'#'*100}")

    buffer_start = bulan_start - pd.Timedelta(days=10)
    sub_calon = panel[(panel["date"] >= buffer_start) &
                       (panel["date"] <= bulan_end)].copy()

    m_a_s = rule_a(sub_calon).fillna(False).values
    m_t_s = rule_t(sub_calon).fillna(False).values
    m_v_s = rule_v(sub_calon).fillna(False).values
    m_union_s = m_a_s | m_t_s | m_v_s

    signals_calon = sub_calon[m_union_s].copy()
    signals_calon["sig_a"] = m_a_s[m_union_s]
    signals_calon["sig_t"] = m_t_s[m_union_s]
    signals_calon["sig_v"] = m_v_s[m_union_s]

    if not signals_calon.empty:
        signals_calon["tanggal_eksekusi"] = signals_calon["date"].apply(next_trading_day)
        signals_bulan = signals_calon[
            (signals_calon["tanggal_eksekusi"] >= bulan_start) &
            (signals_calon["tanggal_eksekusi"] <= bulan_end)
        ].copy()
    else:
        signals_bulan = pd.DataFrame()

    if not signals_bulan.empty:
        total_pl = 0
        n_exec = 0
        n_win = 0
        n_tp = 0
        n_sl = 0
        n_profit = 0
        n_close = 0
        n_skip = 0
        n_pending = 0

        for _, row in signals_bulan.iterrows():
            ret, status = simulate(row)
            if status in ("TP", "PROFIT", "SL", "CLOSE"):
                n_exec += 1
                total_pl += ret
                if ret > 0:
                    n_win += 1
                if status == "TP":
                    n_tp += 1
                elif status == "PROFIT":
                    n_profit += 1
                elif status == "SL":
                    n_sl += 1
                else:
                    n_close += 1
            elif status == "SKIP_GAP_DOWN":
                n_skip += 1
            else:
                n_pending += 1

        print(f"\n  Periode        : {bulan_label}")
        print(f"  Total sinyal   : {len(signals_bulan)}")
        print(f"  Tereksekusi    : {n_exec}")
        print(f"  ├─ TP          : {n_tp}")
        print(f"  ├─ PROFIT      : {n_profit}")
        print(f"  ├─ SL          : {n_sl}")
        print(f"  └─ CLOSE       : {n_close}")
        print(f"  Skip           : {n_skip}")
        print(f"  Pending        : {n_pending}")

        if n_exec > 0:
            print(f"\n  Win rate       : {n_win}/{n_exec} "
                  f"({n_win/n_exec*100:.1f}%)")
            print(f"  Total P/L      : {total_pl*100:+.2f}%")
            print(f"  Avg per trade  : {total_pl/n_exec*100:+.2f}%")
    else:
        print(f"\n  Belum ada sinyal di bulan {bulan_label}.")

    # Simpan
    print(f"\n{'#'*100}")
    print(f"#  MENYIMPAN EXCEL")
    print(f"{'#'*100}\n")

    os.makedirs("quant/output", exist_ok=True)

    if not kandidat.empty:
        csv_path = f"quant/output/kandidat_union_{pred_date.date()}.csv"
        kandidat[["kode", "close", "value", "ret_1d", "ret_5d",
                  "px_vs_sma20", "rk_turnover", "match"]].to_csv(csv_path, index=False)
        print(f"  -> CSV: {csv_path}")

    simpan_excel(kandidat, signals_bulan, pred_date.date(),
                 "quant/output", bulan_label, bulan_file)

    print(f"\n{'='*100}")
    print(f"SELESAI - {datetime.now().strftime('%H:%M:%S')}")
    print(f"{'='*100}")


if __name__ == "__main__":
    main()